#!/usr/bin/env python3
"""Render a data-descriptor Markdown file into the EBRAINS Word template.

Usage:
    python3 fill_word_template.py data-descriptor.md
    python3 fill_word_template.py data-descriptor.md -o out.docx --template word_template.docx

The Markdown file follows templates/markdown_template.md: a running-header line
(*Title: '<short label>' | version: <version>*), then one heading per template
section. The script opens the Word template, removes every guidance paragraph
and the example table, writes each Markdown section under the matching template
heading, rewrites the running header, and validates the result. It exits with
status 1 and a list of problems when the Markdown is incomplete or the filled
document still carries template text.

Markdown the renderer understands, inside a section:

- paragraphs; a line ending in a backslash or two spaces is a hard line break
- `###` / `####` headings not named after a template section: bold subheadings;
  a `#` or `##` heading that matches no template section is an error
- bullet lists (`-`, `*`) and numbered lists (`1.`), rendered with a hanging
  indent and literal numbers
- a fenced block whose first line starts with `/` is the directory tree: one
  paragraph per entry, indented by depth, name in bold, `[description]` in
  italics; any other fenced block is preformatted text
- pipe tables; `<br>` inside a cell is a line break
- `![caption](path)` on its own line: the image, then the caption in italics
- inline `**bold**`, `*italic*`, `` `code` `` (rendered as plain text),
  `^1,2^` superscript, Unicode superscript digits, `[text](url)`

Requires python-docx (`pip install python-docx`).
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from docx.text.paragraph import Paragraph as DocxParagraph
except ImportError:  # pragma: no cover - reported to the user, not tested
    sys.exit("python-docx is required: pip install python-docx")

# Style names as the EBRAINS template defines them. The guidance style is the
# green italic text the curation team uses for instructions; every paragraph
# in it, and the example table, must be gone from a submitted descriptor.
GUIDANCE_STYLE = "Colorful List - Accent 11"
BODY_STYLE = "1) Body Text HBP"
SPACER_STYLE = "Body Text"
PREFORMATTED_STYLE = "Preformatted Text"
TABLE_STYLE = "Table Grid"
HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3"}

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "word_template.docx"

# Limits stated in the template's guidance text.
TITLE_MAX_CHARS = 110
SUMMARY_MAX_WORDS = 220

PAGE_TEXT_WIDTH = Inches(6.5)
TREE_INDENT_PER_LEVEL = Inches(0.2)
LIST_INDENT = Inches(0.3)
CAPTION_FONT_SIZE = Pt(9)
FORMAT_TABLE_HEADER = ("format", "extension", "software used / file specification")
FORMAT_TABLE_WIDTHS = (Inches(1.8), Inches(0.9), Inches(3.8))

TEMPLATE_PLACEHOLDER_RE = re.compile(r"^\[(Add|List) .*\]$")
RESIDUAL_PLACEHOLDER_RE = re.compile(r"\b(TODO|TBD)\b", re.IGNORECASE)
HEADER_LINE_RE = re.compile(r"^\*?\s*Title:\s*['‘]([^'’]+)['’]\s*\|\s*version:\s*(.+?)\s*\*?$")
SUPERSCRIPT_MAP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻", "0123456789+-")


# --------------------------------------------------------------------------
# Markdown model
# --------------------------------------------------------------------------


@dataclass
class ParagraphBlock:
    lines: list[str]  # one entry per hard-broken line


@dataclass
class SubheadingBlock:
    text: str
    level: int


@dataclass
class ListBlock:
    ordered: bool
    items: list[str]


@dataclass
class FenceBlock:
    lines: list[str]


@dataclass
class TableBlock:
    header: list[str]
    rows: list[list[str]]


@dataclass
class ImageBlock:
    caption: str
    path: str


Block = ParagraphBlock | SubheadingBlock | ListBlock | FenceBlock | TableBlock | ImageBlock


@dataclass
class Section:
    name: str
    blocks: list[Block] = field(default_factory=list)


@dataclass
class Descriptor:
    short_title: str
    version: str
    sections: dict[str, Section]  # keyed by normalised heading


def normalize_heading(text: str) -> str:
    return text.strip().rstrip(":").strip().casefold()


# --------------------------------------------------------------------------
# Markdown parsing
# --------------------------------------------------------------------------

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
IMAGE_RE = re.compile(r"^!\[(.*)\]\((.+)\)\s*$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
HARD_BREAK_RE = re.compile(r"(\\|  )$")


def parse_markdown(text: str, section_names: set[str]) -> Descriptor:
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.splitlines()

    short_title = version = None
    sections: dict[str, Section] = {}
    current: Section | None = None
    problems: list[str] = []

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if short_title is None and HEADER_LINE_RE.match(stripped):
            match = HEADER_LINE_RE.match(stripped)
            short_title, version = match.group(1), match.group(2)
            i += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level, title = len(heading.group(1)), heading.group(2)
            key = normalize_heading(title)
            if key in section_names:
                if key in sections:
                    problems.append(f"section '{title}' appears twice")
                current = Section(name=title)
                sections[key] = current
            elif level <= 2:
                problems.append(f"heading '{title}' matches no template section")
            elif current is None:
                problems.append(f"heading '{title}' precedes the first template section")
            else:
                current.blocks.append(SubheadingBlock(text=title, level=level))
            i += 1
            continue

        if current is None:
            problems.append(f"text before the first section heading: '{stripped[:60]}'")
            i += 1
            continue

        if stripped.startswith("```"):
            fence: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                fence.append(lines[i].rstrip())
                i += 1
            current.blocks.append(FenceBlock(lines=fence))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            current.blocks.append(TableBlock(header=rows[0], rows=rows[2:]))
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            current.blocks.append(ImageBlock(caption=image.group(1), path=image.group(2)))
            i += 1
            continue

        list_match = BULLET_RE.match(line) or ORDERED_RE.match(line)
        if list_match:
            ordered = ORDERED_RE.match(line) is not None
            item_re = ORDERED_RE if ordered else BULLET_RE
            items: list[str] = []
            while i < len(lines) and lines[i].strip():
                item = item_re.match(lines[i])
                if item:
                    items.append(item.group(1).strip())
                elif items and lines[i].startswith((" ", "\t")):
                    items[-1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            current.blocks.append(ListBlock(ordered=ordered, items=items))
            continue

        # Paragraph: consecutive non-blank lines that are not another block.
        hard_lines: list[str] = []
        soft: list[str] = []
        while i < len(lines) and lines[i].strip() and not starts_block(lines, i):
            raw = lines[i].rstrip("\n")
            is_break = HARD_BREAK_RE.search(raw) is not None
            soft.append(HARD_BREAK_RE.sub("", raw).strip())
            i += 1
            if is_break:
                hard_lines.append(" ".join(soft))
                soft = []
        if soft:
            hard_lines.append(" ".join(soft))
        current.blocks.append(ParagraphBlock(lines=hard_lines))

    if short_title is None:
        problems.append("missing running-header line: *Title: '<short label>' | version: <version>*")
    if problems:
        raise ValueError("\n".join(problems))
    return Descriptor(short_title=short_title, version=version, sections=sections)


def starts_block(lines: list[str], i: int) -> bool:
    line = lines[i]
    stripped = line.strip()
    if HEADING_RE.match(line) or stripped.startswith("```") or IMAGE_RE.match(stripped):
        return True
    if BULLET_RE.match(line) or ORDERED_RE.match(line):
        return True
    return stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]) is not None


def split_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


# --------------------------------------------------------------------------
# Inline rendering
# --------------------------------------------------------------------------

INLINE_RE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|(?<![\w*])\*(?!\s)(?P<italic>.+?)(?<!\s)\*(?![\w*])"
    r"|(?<!\w)_(?!\s)(?P<italic2>.+?)(?<!\s)_(?!\w)"
    r"|`(?P<code>[^`]+)`"
    r"|\^(?P<sup>[^\s^][^^]*?)\^"
    r"|\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)]+)\)"
    r"|(?P<usup>[⁰¹²³⁴-⁹⁺⁻]+)"
)


def add_inline(paragraph, text: str, *, bold=False, italic=False, superscript=False, size=None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        add_run(paragraph, text[pos:match.start()], bold, italic, superscript, size)
        kind = match.lastgroup
        if kind == "bold":
            add_inline(paragraph, match.group("bold"), bold=True, italic=italic, superscript=superscript, size=size)
        elif kind in ("italic", "italic2"):
            add_inline(paragraph, match.group(kind), bold=bold, italic=True, superscript=superscript, size=size)
        elif kind == "code":
            add_run(paragraph, match.group("code"), bold, italic, superscript, size)
        elif kind == "sup":
            add_inline(paragraph, match.group("sup"), bold=bold, italic=italic, superscript=True, size=size)
        elif kind == "usup":
            add_run(paragraph, match.group("usup").translate(SUPERSCRIPT_MAP), bold, italic, True, size)
        elif kind == "ltext":
            link_text, url = match.group("ltext"), match.group("lurl")
            shown = link_text if link_text.strip() == url.strip() else f"{link_text} ({url})"
            add_run(paragraph, shown, bold, italic, superscript, size)
        pos = match.end()
    add_run(paragraph, text[pos:], bold, italic, superscript, size)


def add_run(paragraph, text: str, bold, italic, superscript, size):
    if not text:
        return
    run = paragraph.add_run(text)
    if bold:
        run.font.bold = True
    if italic:
        run.font.italic = True
    if superscript:
        run.font.superscript = True
    if size is not None:
        run.font.size = size


def add_cell_text(paragraph, text: str) -> None:
    """Render a table cell; `<br>` becomes a line break inside the paragraph."""
    pieces = re.split(r"<br\s*/?>", text, flags=re.IGNORECASE)
    for index, piece in enumerate(pieces):
        if index:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        add_inline(paragraph, piece.strip())


# --------------------------------------------------------------------------
# Block rendering
# --------------------------------------------------------------------------


class Renderer:
    """Writes blocks into the document body in front of an anchor element."""

    def __init__(self, doc, markdown_dir: Path):
        self.doc = doc
        self.markdown_dir = markdown_dir

    def paragraph(self, anchor, style: str):
        paragraph = self.doc.add_paragraph(style=style)
        anchor.addprevious(paragraph._p)
        return paragraph

    def spacer(self, anchor):
        self.paragraph(anchor, SPACER_STYLE)

    def render_section(self, section: Section | None, anchor) -> None:
        if section is None or not section.blocks:
            self.spacer(anchor)
            return
        for block in section.blocks:
            self.render_block(block, anchor)

    def render_block(self, block: Block, anchor) -> None:
        if isinstance(block, ParagraphBlock):
            paragraph = self.paragraph(anchor, BODY_STYLE)
            for index, line in enumerate(block.lines):
                if index:
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                add_inline(paragraph, line)
            self.spacer(anchor)
        elif isinstance(block, SubheadingBlock):
            paragraph = self.paragraph(anchor, BODY_STYLE)
            add_inline(paragraph, block.text, bold=True, italic=block.level >= 4)
            self.spacer(anchor)
        elif isinstance(block, ListBlock):
            for number, item in enumerate(block.items, start=1):
                paragraph = self.paragraph(anchor, BODY_STYLE)
                paragraph.paragraph_format.left_indent = LIST_INDENT
                paragraph.paragraph_format.first_line_indent = -LIST_INDENT
                marker = f"{number}." if block.ordered else "•"
                paragraph.add_run(f"{marker}\t")
                add_inline(paragraph, item)
            self.spacer(anchor)
        elif isinstance(block, FenceBlock):
            if block.lines and block.lines[0].lstrip().startswith("/"):
                self.render_tree(block, anchor)
            else:
                for line in block.lines:
                    paragraph = self.paragraph(anchor, PREFORMATTED_STYLE)
                    paragraph.add_run(line)
            self.spacer(anchor)
        elif isinstance(block, TableBlock):
            self.render_table(block, anchor)
            self.spacer(anchor)
        elif isinstance(block, ImageBlock):
            self.render_image(block, anchor)
            self.spacer(anchor)
        else:  # pragma: no cover - the dataclasses above are the only blocks
            raise TypeError(f"unknown block {block!r}")

    def render_tree(self, block: FenceBlock, anchor) -> None:
        base_indent = min(len(line) - len(line.lstrip(" ")) for line in block.lines if line.strip())
        for line in block.lines:
            if not line.strip():
                continue
            indent = len(line) - len(line.lstrip(" ")) - base_indent
            match = re.match(r"^(.*?)(?:\s+(\[.*\]))?\s*$", line.strip())
            name, description = match.group(1), match.group(2)
            paragraph = self.paragraph(anchor, BODY_STYLE)
            paragraph.paragraph_format.left_indent = TREE_INDENT_PER_LEVEL * (indent // 2)
            paragraph.add_run(name).font.bold = True
            if description:
                paragraph.add_run(" ")
                add_inline(paragraph, description, italic=True)

    def render_table(self, block: TableBlock, anchor) -> None:
        columns = len(block.header)
        table = self.doc.add_table(rows=1 + len(block.rows), cols=columns)
        table.style = TABLE_STYLE
        table.autofit = False
        anchor.addprevious(table._tbl)

        widths = self.column_widths(block)
        for row_index, cells in enumerate([block.header] + block.rows):
            row = table.rows[row_index]
            for col_index in range(columns):
                cell = row.cells[col_index]
                cell.width = widths[col_index]
                paragraph = cell.paragraphs[0]
                paragraph.style = self.doc.styles[SPACER_STYLE]
                text = cells[col_index] if col_index < len(cells) else ""
                if row_index == 0:
                    add_inline(paragraph, text, bold=True)
                else:
                    add_cell_text(paragraph, text)

    @staticmethod
    def column_widths(block: TableBlock):
        header = tuple(h.strip().casefold() for h in block.header)
        if header == FORMAT_TABLE_HEADER:
            return FORMAT_TABLE_WIDTHS
        columns = len(block.header)
        return tuple(Inches(6.5 / columns) for _ in range(columns))

    def render_image(self, block: ImageBlock, anchor) -> None:
        path = (self.markdown_dir / block.path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"image not found: {block.path} (resolved to {path})")
        self.doc.add_picture(str(path))
        shape = self.doc.inline_shapes[-1]
        if shape.width > PAGE_TEXT_WIDTH:
            ratio = PAGE_TEXT_WIDTH / shape.width
            shape.height = int(shape.height * ratio)
            shape.width = PAGE_TEXT_WIDTH
        picture_paragraph = self.doc.paragraphs[-1]
        anchor.addprevious(picture_paragraph._p)
        if block.caption.strip():
            caption = self.paragraph(anchor, BODY_STYLE)
            add_inline(caption, block.caption.strip(), italic=True, size=CAPTION_FONT_SIZE)


# --------------------------------------------------------------------------
# Template filling
# --------------------------------------------------------------------------


def template_sections(doc) -> list[tuple[str, object]]:
    """Return (normalised heading, heading element) for each template heading."""
    found = []
    for element in doc.element.body.iterchildren():
        if element.tag != qn("w:p"):
            continue
        paragraph = DocxParagraph(element, doc._body)
        if paragraph.style is not None and paragraph.style.name in HEADING_STYLES:
            found.append((normalize_heading(paragraph.text), element))
    return found


def clear_template_body(doc) -> None:
    """Remove everything from the body except headings and the section properties."""
    body = doc.element.body
    for element in list(body.iterchildren()):
        if element.tag == qn("w:sectPr"):
            continue
        if element.tag == qn("w:p"):
            paragraph = DocxParagraph(element, doc._body)
            if paragraph.style is not None and paragraph.style.name in HEADING_STYLES:
                continue
        body.remove(element)


def fill_document(doc, descriptor: Descriptor, markdown_dir: Path) -> None:
    sections = template_sections(doc)
    clear_template_body(doc)
    renderer = Renderer(doc, markdown_dir)
    body = doc.element.body
    for index, (key, _heading) in enumerate(sections):
        anchor = sections[index + 1][1] if index + 1 < len(sections) else body.sectPr
        renderer.render_section(descriptor.sections.get(key), anchor)
    set_running_header(doc, descriptor)


def set_running_header(doc, descriptor: Descriptor) -> None:
    new_text = f"Title: ‘{descriptor.short_title}’ | version: {descriptor.version}"
    replaced = False
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "version:" not in paragraph.text:
                continue
            runs = [run for run in paragraph.runs if run.text]
            runs[0].text = new_text
            for run in runs[1:]:
                run.text = ""
            replaced = True
    if not replaced:
        raise ValueError("the template header has no 'version:' paragraph to rewrite")


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------


def check_markdown(descriptor: Descriptor, section_keys: list[str]) -> list[str]:
    problems = []
    content_keys = [k for k in section_keys if k != "data descriptor"]
    missing = [k for k in content_keys if k not in descriptor.sections or not descriptor.sections[k].blocks]
    if missing:
        problems.append("sections missing or empty: " + ", ".join(missing))

    title = section_text(descriptor, "title")
    if title:
        if len(title) > TITLE_MAX_CHARS:
            problems.append(f"title is {len(title)} characters; the template allows {TITLE_MAX_CHARS}")
        if re.search(r"[:()]", title):
            problems.append("title contains a colon or parenthesis, which the template forbids")

    summary = section_text(descriptor, "summary")
    if summary:
        words = len(summary.split())
        if words > SUMMARY_MAX_WORDS:
            problems.append(f"summary is {words} words; the template allows {SUMMARY_MAX_WORDS}")
        if re.search(r"\[\d+\]|\^\d", summary):
            problems.append("summary contains a reference citation, which the template forbids")

    for key, section in descriptor.sections.items():
        for block in section.blocks:
            for text in block_texts(block):
                if RESIDUAL_PLACEHOLDER_RE.search(text):
                    problems.append(f"placeholder text in section '{section.name}': '{text[:60]}'")
    return problems


def section_text(descriptor: Descriptor, key: str) -> str:
    section = descriptor.sections.get(key)
    if section is None:
        return ""
    return " ".join(t for block in section.blocks for t in block_texts(block))


def block_texts(block: Block) -> list[str]:
    if isinstance(block, ParagraphBlock):
        return block.lines
    if isinstance(block, SubheadingBlock):
        return [block.text]
    if isinstance(block, ListBlock):
        return block.items
    if isinstance(block, FenceBlock):
        return block.lines
    if isinstance(block, TableBlock):
        return [c for row in [block.header] + block.rows for c in row]
    if isinstance(block, ImageBlock):
        return [block.caption]
    return []


def check_document(doc) -> list[str]:
    """Find template text that survived filling."""
    problems = []
    for paragraph in iter_all_paragraphs(doc):
        style = paragraph.style.name if paragraph.style is not None else ""
        if style == GUIDANCE_STYLE:
            problems.append(f"guidance paragraph survived: '{paragraph.text[:60]}'")
        if TEMPLATE_PLACEHOLDER_RE.match(paragraph.text.strip()):
            problems.append(f"template placeholder survived: '{paragraph.text.strip()}'")
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "Short Title" in paragraph.text or "short name of the DSV" in paragraph.text:
                problems.append("running header still holds the template text")
    return problems


def iter_all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def render_descriptor(markdown_path: Path, output_path: Path, template_path: Path) -> list[str]:
    """Fill the template from the Markdown file. Returns the list of problems."""
    doc = Document(str(template_path))
    section_keys = [key for key, _ in template_sections(doc)]
    descriptor = parse_markdown(markdown_path.read_text(encoding="utf-8"), set(section_keys))

    problems = check_markdown(descriptor, section_keys)
    fill_document(doc, descriptor, markdown_path.parent)
    problems += check_document(doc)
    doc.save(str(output_path))
    return problems


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("markdown", type=Path, help="data-descriptor Markdown file")
    parser.add_argument("-o", "--output", type=Path, help="output .docx (default: beside the Markdown file)")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="EBRAINS Word template")
    args = parser.parse_args(argv)

    output = args.output or args.markdown.with_suffix(".docx")
    try:
        problems = render_descriptor(args.markdown, output, args.template)
    except (ValueError, FileNotFoundError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    if problems:
        print(f"wrote {output}, but the descriptor is not ready:", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        return 1
    print(f"wrote {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
